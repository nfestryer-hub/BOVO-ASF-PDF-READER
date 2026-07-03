"""Module métier des inspections ASF (ITV) de Bovo ASF Manager.

Gère les modules éditables issus du classeur Excel — Inspections, Rapports
d'inspection, Incidents et Matériel — entièrement dans le logiciel (plus besoin
d'Excel pour la saisie). Fournit également :

    * l'import initial du classeur Excel (seeding) ;
    * les listes de référence alimentant les menus déroulants ;
    * le recalcul des indicateurs (KPIs) du tableau de bord ;
    * la complétion automatique des interventions : dès qu'un rapport existe
      pour un ouvrage (champ « Identifiant »), l'intervention correspondante
      passe au statut « Terminée ».

Aucune logique d'affichage ici : le GUI consomme uniquement ce service.
"""

from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import openpyxl

from app.database import Database
from app.logger import LogAction, LogLevel, get_logger, log_action
from app.models import StatutIntervention

logger = get_logger("inspections")


@dataclass(slots=True, frozen=True)
class FieldSpec:
    """Description d'un champ de formulaire pour un module."""

    column: str
    label: str
    widget: str = "text"  # text | number | date | long | choice
    ref_key: str = ""


# Colonnes des tables, dans l'ordre exact des colonnes du classeur Excel.
ORDERED_COLUMNS: dict[str, list[str]] = {
    "inspections": [
        "n_inspection", "date", "jour_nuit", "autoroute", "secteur",
        "pk_depart", "pk_arrivee", "ml_inspectes", "nb_regards",
        "heure_debut", "heure_fin", "heures_totales", "ml_h",
        "operateur", "coequipier", "meteo", "observations",
    ],
    "rapports": [
        "date", "identifiant", "n_inspection", "prestataire", "inspecteur",
        "autoroute", "type_ouvrage", "secteur", "sens", "emplacement",
        "type_materiau", "latitude", "longitude", "lien_maps", "requalif",
        "etat_cunettes", "type_defaut", "urgence", "description", "photo",
        "statut",
    ],
    "incidents": [
        "date", "autoroute", "pk", "type_incident", "description",
        "impact", "action_prise",
    ],
    "materiel": [
        "date", "camera", "chariot", "vehicule", "km_depart", "km_arrivee",
        "entretien", "panne", "observations",
    ],
}

SHEET_LABELS: dict[str, str] = {
    "inspections": "Inspections",
    "rapports": "Rapports d'inspection",
    "incidents": "Incidents",
    "materiel": "Matériel",
}

# Noms des feuilles dans le classeur Excel source.
EXCEL_SHEET_NAMES: dict[str, str] = {
    "inspections": "Inspections",
    "rapports": "Rapport d'Inspections",
    "incidents": "Incidents",
    "materiel": "Matériel",
}

FORM_FIELDS: dict[str, list[FieldSpec]] = {
    "inspections": [
        FieldSpec("n_inspection", "N° Inspection", "text"),
        FieldSpec("date", "Date", "date"),
        FieldSpec("jour_nuit", "Jour/Nuit", "choice", "choix_jour"),
        FieldSpec("autoroute", "Autoroute", "choice", "autoroutes"),
        FieldSpec("secteur", "Secteur", "choice", "secteur"),
        FieldSpec("pk_depart", "PK Départ", "text"),
        FieldSpec("pk_arrivee", "PK Arrivée", "text"),
        FieldSpec("ml_inspectes", "ML Inspectés", "number"),
        FieldSpec("nb_regards", "Nb Regards", "number"),
        FieldSpec("heure_debut", "Heure Début", "text"),
        FieldSpec("heure_fin", "Heure Fin", "text"),
        FieldSpec("heures_totales", "Heures Totales", "number"),
        FieldSpec("operateur", "Opérateur", "choice", "operateurs_itv"),
        FieldSpec("coequipier", "Coéquipier", "choice", "operateurs_hydro"),
        FieldSpec("meteo", "Météo", "choice", "meteo"),
        FieldSpec("observations", "Observations", "long"),
    ],
    "rapports": [
        FieldSpec("date", "Date de visite", "date"),
        FieldSpec("identifiant", "Identifiant", "text"),
        FieldSpec("n_inspection", "N° Inspection", "text"),
        FieldSpec("prestataire", "Prestataire", "text"),
        FieldSpec("inspecteur", "Inspecteur", "choice", "operateurs_itv"),
        FieldSpec("autoroute", "Autoroute", "choice", "autoroutes"),
        FieldSpec("type_ouvrage", "Type d'ouvrage", "choice", "type"),
        FieldSpec("secteur", "Secteur", "choice", "secteur"),
        FieldSpec("sens", "Sens", "text"),
        FieldSpec("emplacement", "Emplacement", "choice", "emplacement"),
        FieldSpec("type_materiau", "Type de matériau", "choice", "type_materiau"),
        FieldSpec("latitude", "Latitude", "text"),
        FieldSpec("longitude", "Longitude", "text"),
        FieldSpec("lien_maps", "Lien Maps", "text"),
        FieldSpec("requalif", "Si objet à requalifier", "choice", "requalif"),
        FieldSpec("etat_cunettes", "État général de l'ouvrage", "choice", "etat_cunettes"),
        FieldSpec("type_defaut", "Type de défaut", "text"),
        FieldSpec("urgence", "Urgence", "choice", "urgence"),
        FieldSpec("description", "Commentaires", "long"),
        FieldSpec("photo", "Photo(s)", "text"),
        FieldSpec("statut", "Statut", "text"),
    ],
    "incidents": [
        FieldSpec("date", "Date", "date"),
        FieldSpec("autoroute", "Autoroute", "choice", "autoroutes"),
        FieldSpec("pk", "PK", "text"),
        FieldSpec("type_incident", "Type Incident", "text"),
        FieldSpec("description", "Description", "long"),
        FieldSpec("impact", "Impact", "text"),
        FieldSpec("action_prise", "Action prise", "long"),
    ],
    "materiel": [
        FieldSpec("date", "Date", "date"),
        FieldSpec("camera", "Caméra", "choice", "camera"),
        FieldSpec("chariot", "Chariot", "choice", "chariot"),
        FieldSpec("vehicule", "Véhicule", "choice", "vehicules"),
        FieldSpec("km_depart", "Km Départ", "text"),
        FieldSpec("km_arrivee", "Km Arrivée", "text"),
        FieldSpec("entretien", "Entretien", "text"),
        FieldSpec("panne", "Panne", "choice", "panne"),
        FieldSpec("observations", "Observations", "long"),
    ],
}

# Colonnes affichées dans les tableaux récapitulatifs de chaque module.
LIST_COLUMNS: dict[str, list[str]] = {
    "inspections": ["n_inspection", "date", "autoroute", "secteur", "ml_inspectes", "operateur"],
    "rapports": ["date", "identifiant", "type_ouvrage", "autoroute", "secteur", "etat_cunettes", "urgence"],
    "incidents": ["date", "autoroute", "pk", "type_incident", "impact"],
    "materiel": ["date", "camera", "vehicule", "panne", "entretien"],
}

# Découpage du formulaire "Rapport d'inspection" en 4 sections ASF
# (basé sur le classeur ASF "Form visites assain V2026-06-01.xlsx").
# Chaque section = (titre, [(colonne_db, libellé)]).
ASF_RAPPORT_SECTIONS: list[tuple[str, list[tuple[str, str]]]] = [
    (
        "Informations Générales",
        [
            ("identifiant", "Identifiant"),
            ("date", "Date de visite"),
            ("prestataire", "Prestataire"),
            ("inspecteur", "Inspecteur"),
            ("n_inspection", "N° Inspection"),
            ("autoroute", "Autoroute"),
            ("secteur", "Secteur"),
            ("sens", "Sens"),
        ],
    ),
    (
        "Caractéristiques de l'ouvrage",
        [
            ("type_ouvrage", "Type d'ouvrage"),
            ("type_materiau", "Type de matériau"),
            ("emplacement", "Emplacement"),
            ("latitude", "Latitude"),
            ("longitude", "Longitude"),
            ("lien_maps", "Lien Maps"),
        ],
    ),
    (
        "Surveillance générale",
        [
            ("requalif", "Si objet à requalifier"),
            ("etat_cunettes", "État général de l'ouvrage"),
            ("type_defaut", "Type de défaut"),
            ("urgence", "Urgence"),
            ("statut", "Statut"),
        ],
    ),
    (
        "Autres documents rattachés à la visite",
        [
            ("description", "Commentaires"),
            ("photo", "Photo(s)"),
        ],
    ),
]

# Listes de référence par défaut (utilisées tant que le classeur n'est pas importé).
DEFAULT_REFERENTIEL: dict[str, list[str]] = {
    "operateurs_itv": ["Hugo D.", "Dylan H."],
    "operateurs_hydro": [
        "Ludovic B.", "Amadou B.", "Dylan P.", "Exaucé", "Pierre",
        "Tristan", "Yassine", "Baptiste", "Yohan",
    ],
    "meteo": ["Beau", "Pluvieux", "Nuageux", "Orageux", "Neigeux"],
    "choix_jour": ["Jour", "Nuit"],
    "secteur": ["Midi Toulousain", "Garonne", "Agen", "Quercy"],
    "autoroutes": ["A61", "A66", "A68", "A62", "A20", "A64"],
    # Types d'ouvrage ASF (formulaire visites assainissement V2026-06-01).
    "type": [
        "Caniveau", "Collecteur enterré", "Cunette", "Fossé",
        "Descente d'eau", "CA", "OH", "DE",
    ],
    "type_materiau": [
        "Enherbé", "Métallique", "Mixte", "Béton",
        "Cunette revêtue - béton", "Cunette revêtue - autre matériau",
        "Cunette non revêtue",
        "Fossé revêtu – béton", "Fossé revêtu – autre matériau",
        "Fossé non revêtu",
        "Descente d'eau enterrée", "Descente d'eau superficielle",
        "Caniveau à fente", "Caniveau à grille",
        "Caniveau ouvert", "Caniveau recouvert",
        "Collecteur enterré",
        "Autre",
    ],
    "emplacement": [
        "Haut de talus", "Bande dérasée", "Derrière GBA",
        "Derrière glissière", "Bas de talus", "Autre",
    ],
    "urgence": ["Aucune", "Moindre", "Moyenne", "Elevée", "Très elevée"],
    "requalif": [
        "Non", "Caniveau à requalifier", "Cunette à requalifier",
        "Fossé à requalifier", "Descente d'eau à requalifier",
        "Collecteur enterré à requalifier",
    ],
    "etat_cunettes": [
        "Pas de défaut", "Non évolutif", "Evolution normale",
        "Défauts majeurs", "Non visitable",
    ],
    "camera": ["IBAK Orpheus 2", "IBAK Nano", "IBAK Nano L", "IBAK Orion 3"],
    "chariot": ["T76.2", "T66.1"],
    "vehicules": ["Sprinter 41", "Citan 55"],
    "panne": ["Aucune", "Oui"],
}

# Index des colonnes du Référentiel Excel → clé de référence.
_REFERENTIEL_COLUMN_MAP: dict[int, str] = {
    1: "operateurs_itv", 2: "operateurs_hydro", 4: "meteo", 6: "choix_jour",
    7: "secteur", 8: "autoroutes", 9: "type", 10: "type_materiau",
    11: "urgence", 12: "requalif", 13: "etat_cunettes", 16: "camera",
    17: "chariot", 18: "vehicules", 19: "panne",
}


def _cell_to_str(value: Any) -> str:
    """Convertit une valeur de cellule Excel en chaîne stockable."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, float):
        return str(round(value, 3))
    return str(value).strip()


def _to_float(value: str) -> float:
    """Convertit une chaîne en flottant, ou 0.0 si impossible."""
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return 0.0


class InspectionService:
    """Service applicatif des modules ASF (CRUD, import Excel, KPIs, liaisons)."""

    SHEETS: tuple[str, ...] = ("inspections", "rapports", "incidents", "materiel")

    def __init__(self, database: Database, audit: Any | None = None) -> None:
        """Initialise le service et garantit l'existence des tables.

        Args:
            database: Instance :class:`Database` partagée.
            audit: :class:`AuditService` optionnel pour enregistrer l'historique.
        """
        self._db = database
        self._db.ensure_tables_from_specs(ORDERED_COLUMNS)
        self._audit = audit
        self._referentiel = self._load_referentiel()

    # ------------------------------------------------------------------ #
    # Référentiel
    # ------------------------------------------------------------------ #
    def _load_referentiel(self) -> dict[str, list[str]]:
        """Charge le référentiel depuis les paramètres, sinon les défauts."""
        stored = self._db.get_setting("referentiel")
        if stored:
            try:
                data = json.loads(stored)
                merged = dict(DEFAULT_REFERENTIEL)
                merged.update({k: v for k, v in data.items() if v})
                return merged
            except json.JSONDecodeError:
                logger.warning("Référentiel stocké illisible, utilisation des défauts.")
        return dict(DEFAULT_REFERENTIEL)

    def referentiel(self, ref_key: str) -> list[str]:
        """Retourne la liste de référence associée à une clé."""
        return self._referentiel.get(ref_key, [])

    # ------------------------------------------------------------------ #
    # Métadonnées de structure (consommées par le GUI)
    # ------------------------------------------------------------------ #
    @staticmethod
    def sheet_label(sheet: str) -> str:
        """Libellé lisible d'un module."""
        return SHEET_LABELS[sheet]

    @staticmethod
    def form_fields(sheet: str) -> list[FieldSpec]:
        """Champs de formulaire d'un module."""
        return FORM_FIELDS[sheet]

    @staticmethod
    def list_columns(sheet: str) -> list[tuple[str, str]]:
        """Colonnes (clé, libellé) affichées dans le tableau d'un module."""
        labels = {f.column: f.label for f in FORM_FIELDS[sheet]}
        return [(col, labels.get(col, col)) for col in LIST_COLUMNS[sheet]]

    # ------------------------------------------------------------------ #
    # CRUD
    # ------------------------------------------------------------------ #
    def list_rows(self, sheet: str, search: str = "") -> list[dict[str, str]]:
        """Liste les lignes d'un module."""
        return self._db.list_generic(sheet, search)

    def get_row(self, sheet: str, row_id: int) -> dict[str, str] | None:
        """Retourne une ligne d'un module par son id."""
        return self._db.get_generic(sheet, row_id)

    def save_row(
        self, sheet: str, data: dict[str, str], row_id: int | None = None
    ) -> int:
        """Crée ou met à jour une ligne, puis déclenche les liaisons éventuelles.

        Returns:
            L'identifiant de la ligne enregistrée.
        """
        before: dict[str, str] = {}
        if row_id is None:
            new_id = self._db.insert_generic(sheet, data)
            log_action(
                self._db, LogAction.CREATE,
                f"{SHEET_LABELS[sheet]} : nouvelle entrée #{new_id}.",
            )
            self._audit_record("create", sheet, str(new_id), {}, data)
        else:
            before = self._db.get_generic(sheet, row_id) or {}
            self._db.update_generic(sheet, row_id, data)
            new_id = row_id
            log_action(
                self._db, LogAction.UPDATE,
                f"{SHEET_LABELS[sheet]} : entrée #{row_id} modifiée.",
            )
            self._audit_record("update", sheet, str(row_id), before, data)
        if sheet == "rapports":
            self._mark_intervention_done(data.get("identifiant", ""))
        return new_id

    def delete_row(self, sheet: str, row_id: int) -> None:
        """Supprime une ligne d'un module."""
        before = self._db.get_generic(sheet, row_id) or {}
        self._db.delete_generic(sheet, row_id)
        log_action(
            self._db, LogAction.DELETE,
            f"{SHEET_LABELS[sheet]} : entrée #{row_id} supprimée.",
            level=LogLevel.WARNING,
        )
        self._audit_record("delete", sheet, str(row_id), before, {})

    def _audit_record(
        self, action: str, entity: str, entity_id: str,
        before: dict[str, Any], after: dict[str, Any],
    ) -> None:
        """Enregistre une modification dans l'audit trail si le service est branché."""
        if self._audit is None:
            return
        try:
            self._audit.record(
                action=action, entity=entity, entity_id=entity_id,
                before=before, after=after,
            )
        except Exception as exc:  # noqa: BLE001 - audit best-effort
            logger.warning("Audit non enregistré : %s", exc)

    # ------------------------------------------------------------------ #
    # Liaison rapport → intervention
    # ------------------------------------------------------------------ #
    def _mark_intervention_done(self, identifiant: str) -> int:
        """Passe l'intervention liée au statut « Terminée ». Retourne 1 si traité."""
        identifiant = (identifiant or "").strip()
        if not identifiant:
            return 0
        intervention = self._db.get_intervention_by_ouvrage(identifiant)
        if intervention is None or intervention.id is None:
            return 0
        if intervention.statut is StatutIntervention.TERMINEE:
            return 0
        self._db.update_intervention_fields(
            intervention.id, statut=StatutIntervention.TERMINEE.value
        )
        log_action(
            self._db, LogAction.UPDATE,
            f"Intervention {identifiant} marquée « Terminée » (rapport ASF).",
            level=LogLevel.SUCCESS,
        )
        return 1

    def synchronize_interventions(self) -> int:
        """Re-synchronise toutes les interventions à partir des rapports existants."""
        count = 0
        for rapport in self._db.list_generic("rapports"):
            count += self._mark_intervention_done(rapport.get("identifiant", ""))
        return count

    # ------------------------------------------------------------------ #
    # Pré-remplissage du rapport à partir de la dernière inspection
    # ------------------------------------------------------------------ #
    #: Correspondance ``colonne_rapport → colonne_inspection`` reprise
    #: automatiquement quand on crée un nouveau rapport d'inspection.
    RAPPORT_PREFILL_MAP: dict[str, str] = {
        "date": "date",
        "n_inspection": "n_inspection",
        "autoroute": "autoroute",
        "secteur": "secteur",
        "inspecteur": "operateur",
    }

    #: Colonnes du rapport laissées vides volontairement (à saisir).
    RAPPORT_REQUIRED_MANUAL: tuple[str, ...] = (
        "identifiant", "type_ouvrage", "sens", "emplacement", "type_materiau",
        "latitude", "longitude", "lien_maps", "requalif", "etat_cunettes",
        "type_defaut", "urgence", "description", "photo", "statut",
    )

    def latest_inspection(self) -> dict[str, str] | None:
        """Retourne l'inspection la plus récemment ajoutée, ou ``None``."""
        rows = self._db.list_generic("inspections")
        return rows[0] if rows else None

    def prefill_rapport(self) -> dict[str, str]:
        """Retourne les valeurs à pré-remplir dans un nouveau rapport d'inspection.

        Puise dans la dernière inspection ajoutée. Ajoute ``prestataire =
        "BOVO & FILS"`` par défaut. Les colonnes de :attr:`RAPPORT_REQUIRED_MANUAL`
        restent vides pour forcer la saisie utilisateur.
        """
        prefill: dict[str, str] = {"prestataire": "BOVO & FILS"}
        inspection = self.latest_inspection()
        if inspection is None:
            return prefill
        for rapport_col, inspection_col in self.RAPPORT_PREFILL_MAP.items():
            value = (inspection.get(inspection_col) or "").strip()
            if value:
                prefill[rapport_col] = value
        return prefill

    # ------------------------------------------------------------------ #
    # KPIs (recalcul du tableau de bord Excel)
    # ------------------------------------------------------------------ #
    def kpis(self) -> dict[str, str]:
        """Recalcule les indicateurs ASF à partir des inspections."""
        rows = self._db.list_generic("inspections")
        total_ml = sum(_to_float(r.get("ml_inspectes", "")) for r in rows)
        total_regards = sum(_to_float(r.get("nb_regards", "")) for r in rows)
        total_heures = sum(_to_float(r.get("heures_totales", "")) for r in rows)
        ml_h = total_ml / total_heures if total_heures else 0.0
        return {
            "ml_inspectes": f"{total_ml:,.0f}".replace(",", " "),
            "nb_inspections": str(len(rows)),
            "regards": f"{total_regards:,.0f}".replace(",", " "),
            "heures": f"{total_heures:.1f}",
            "ml_h": f"{ml_h:.1f}",
        }

    # ------------------------------------------------------------------ #
    # Export PDF (mode tableur) d'un module ASF
    # ------------------------------------------------------------------ #
    def export_pdf(self, sheet: str, destination: Path) -> Path:
        """Exporte toutes les lignes d'un module ASF dans un PDF.

        Pour ``sheet == "rapports"``, utilise la mise en page ASF officielle
        (4 sections colorées : Informations Générales / Caractéristiques /
        Surveillance / Autres documents), conforme au classeur
        ``Form visites assain V2026-06-01.xlsx``.

        Pour les autres modules, utilise le mode tableur générique.
        """
        from app.exporter import Exporter  # import local pour éviter le cycle
        exporter = Exporter(self._db)
        rows_data = self._db.list_generic(sheet)
        if sheet == "rapports":
            return exporter.asf_rapports_to_pdf(
                ASF_RAPPORT_SECTIONS, rows_data, destination,
            )
        columns = ORDERED_COLUMNS[sheet]
        labels = {f.column: f.label for f in FORM_FIELDS[sheet]}
        headers = [labels.get(c, c) for c in columns]
        rows = [[(row.get(c, "") or "") for c in columns] for row in rows_data]
        title = f"Bovo ASF Manager — {SHEET_LABELS[sheet]}"
        return exporter.table_to_pdf(title, headers, rows, destination)

    # ------------------------------------------------------------------ #
    # Import du classeur Excel
    # ------------------------------------------------------------------ #
    def import_workbook(self, path: Path) -> dict[str, int]:
        """Importe le classeur ASF dans SQLite (remplace les données existantes).

        Args:
            path: Chemin du fichier ``.xlsx``.

        Returns:
            Un récapitulatif ``{module: nombre_de_lignes}``.

        Raises:
            FileNotFoundError: Si le fichier est introuvable.
        """
        if not path.exists():
            raise FileNotFoundError(f"Classeur introuvable : {path}")
        workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        summary: dict[str, int] = {}
        try:
            self._import_referentiel(workbook)
            for sheet, excel_name in EXCEL_SHEET_NAMES.items():
                if excel_name not in workbook.sheetnames:
                    summary[sheet] = 0
                    continue
                summary[sheet] = self._import_sheet(workbook[excel_name], sheet)
        finally:
            workbook.close()

        linked = self.synchronize_interventions()
        log_action(
            self._db, LogAction.IMPORT,
            f"Classeur ASF importé : {summary} · {linked} intervention(s) liée(s).",
            level=LogLevel.SUCCESS,
        )
        summary["interventions_liees"] = linked
        return summary

    def _import_sheet(self, worksheet: Any, sheet: str) -> int:
        """Importe une feuille Excel dans la table correspondante."""
        columns = ORDERED_COLUMNS[sheet]
        anchor = "date" if "date" in columns else columns[0]
        anchor_index = columns.index(anchor)
        self._db.clear_table(sheet)
        inserted = 0
        for row_index, row in enumerate(worksheet.iter_rows(values_only=True)):
            if row_index == 0:
                continue  # ligne d'en-tête
            values = list(row)
            anchor_value = values[anchor_index] if anchor_index < len(values) else None
            if anchor_value in (None, ""):
                continue  # ligne-modèle vide (pas de date réelle)
            data = {
                columns[i]: _cell_to_str(values[i])
                for i in range(min(len(columns), len(values)))
            }
            self._db.insert_generic(sheet, data)
            inserted += 1
        return inserted

    def _import_referentiel(self, workbook: Any) -> None:
        """Extrait les listes de référence du classeur et les persiste."""
        if "Référentiel" not in workbook.sheetnames:
            return
        worksheet = workbook["Référentiel"]
        collected: dict[str, list[str]] = {key: [] for key in _REFERENTIEL_COLUMN_MAP.values()}
        for row_index, row in enumerate(worksheet.iter_rows(values_only=True)):
            if row_index == 0:
                continue
            for col_index, ref_key in _REFERENTIEL_COLUMN_MAP.items():
                cell = row[col_index - 1] if col_index - 1 < len(row) else None
                text = _cell_to_str(cell)
                if text and text not in collected[ref_key]:
                    collected[ref_key].append(text)
        merged = dict(DEFAULT_REFERENTIEL)
        merged.update({k: v for k, v in collected.items() if v})
        self._referentiel = merged
        self._db.set_setting("referentiel", json.dumps(merged, ensure_ascii=False))


__all__ = ["FieldSpec", "InspectionService"]
